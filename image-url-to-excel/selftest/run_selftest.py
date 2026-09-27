#!/usr/bin/env python3
"""端到端自测：起本地图片服务 → 造样例表 → 跑 url2img.py → 校验结果。

    python selftest/run_selftest.py

服务端故意模拟三种真实世界情况：防盗链校验、扩展名与实际格式不符、404 坏链。
"""

from __future__ import annotations

import csv
import functools
import os
import re
import shutil
import subprocess
import sys
import threading
import zipfile
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from PIL import Image, ImageDraw  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

WORK = Path(__file__).resolve().parent / "_work"
IMG_DIR = WORK / "images"
HITS: dict[str, int] = {}

SIZE = 100
DATA_FIRST, DATA_LAST = 2, 12   # 样例表的数据行范围（第 1 行是表头）
BAD_ROWS = {8}                  # 404 坏链那一行，预期下载失败
FAILURES: list[str] = []


def check(condition: bool, label: str) -> None:
    print(f"  {'PASS' if condition else 'FAIL'}  {label}")
    if not condition:
        FAILURES.append(label)


# --------------------------------------------------------------------------
# 图片素材与本地服务
# --------------------------------------------------------------------------

def make_images() -> None:
    IMG_DIR.mkdir(parents=True, exist_ok=True)

    img = Image.new("RGB", (800, 600))
    draw = ImageDraw.Draw(img)
    for x in range(0, 800, 40):
        draw.rectangle([x, 0, x + 20, 600], fill=(x % 255, 120, (200 - x) % 255))
    img.save(IMG_DIR / "normal.jpg", "JPEG", quality=90)

    img = Image.new("RGBA", (400, 400), (0, 0, 0, 0))
    ImageDraw.Draw(img).ellipse([100, 100, 300, 300], fill=(220, 40, 40, 255))
    img.save(IMG_DIR / "transparent.png")

    # 扩展名写 .jpg，内容其实是 WebP
    Image.new("RGB", (800, 600), (30, 90, 200)).save(
        IMG_DIR / "actual_webp.jpg", "WEBP", quality=90)

    Image.new("RGB", (1600, 400), (250, 200, 60)).save(IMG_DIR / "wide.png")
    Image.new("RGB", (400, 1600), (80, 180, 140)).save(
        IMG_DIR / "tall.jpg", "JPEG", quality=90)
    Image.new("RGB", (600, 600), (140, 90, 200)).save(
        IMG_DIR / "hotlink.jpg", "JPEG", quality=90)


class Handler(SimpleHTTPRequestHandler):
    def do_GET(self):
        path = urlparse(self.path).path
        HITS[path] = HITS.get(path, 0) + 1
        if path == "/hotlink.jpg" and not self.headers.get("Referer"):
            self.send_error(403, "hotlink blocked")
            return
        super().do_GET()

    def log_message(self, *args):
        pass


def start_server() -> tuple[ThreadingHTTPServer, int]:
    server = ThreadingHTTPServer(
        ("127.0.0.1", 0), functools.partial(Handler, directory=str(IMG_DIR)))
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return server, server.server_address[1]


# --------------------------------------------------------------------------
# 样例表格
# --------------------------------------------------------------------------

def make_sample(port: int) -> Path:
    base = f"http://127.0.0.1:{port}"
    rows = [
        ("A 正常 jpg", f"{base}/normal.jpg", 99),
        ("B 透明 png", f"{base}/transparent.png", 88),
        ("C 伪后缀 webp", f"{base}/actual_webp.jpg", 77),
        ("D 宽图", f"{base}/wide.png", 66),
        ("E 长图", f"{base}/tall.jpg", 55),
        ("F 防盗链", f"{base}/hotlink.jpg", 44),
        ("G 坏链", f"{base}/missing.jpg", 33),
        ("H 带参数", f"{base}/normal.jpg?w=200", 22),
        ("I 夹在文字里", f"商品图片：{base}/normal.jpg 请查收", 11),
        ("J 同格多图", f"{base}/normal.jpg,{base}/wide.png", 10),
        # 真实电商后台常见形态：JSON 里塞同一张图的多种尺寸，
        # 且 x-oss-process 的逗号没有转义成 %2C
        ("K JSON 多尺寸",
         '{"pic400x400":"' + base
         + '/normal.jpg?x-oss-process=image/resize,m_mfit,w_400,h_400/format,jpg/quality,Q_85"'
         + ',"pic50x50":"' + base
         + '/normal.jpg?x-oss-process=image/resize,m_mfit,w_50,h_50/format,jpg/quality,Q_85"}',
         5),
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "商品"
    ws.append(["商品名称", "主图链接", "价格"])
    for row in rows:
        ws.append(row)
    path = WORK / "sample.xlsx"
    wb.save(path)

    with (WORK / "sample.csv").open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh)
        writer.writerow(["商品名称", "主图链接", "价格"])
        writer.writerows(rows)
    return path


# --------------------------------------------------------------------------
# 运行与校验
# --------------------------------------------------------------------------

def run_script(*args: str) -> None:
    result = subprocess.run([sys.executable, str(ROOT / "url2img.py"), *args],
                            capture_output=True, text=True, encoding="utf-8",
                            errors="replace", cwd=str(ROOT), env=child_env())
    print(result.stdout.rstrip())
    if result.returncode != 0:
        print(result.stderr.rstrip())
        raise SystemExit(f"脚本退出码 {result.returncode}")


def child_env() -> dict:
    """脚本不再自行修改 stdout 编码，这里显式指定，便于按 UTF-8 读取输出。"""
    return {**os.environ, "PYTHONIOENCODING": "utf-8"}


def check_bat() -> None:
    print("\n[运行 5] 双击运行的 BAT 启动器")
    bat = ROOT / "run.bat"
    if sys.platform != "win32":
        print("  SKIP  非 Windows 平台")
        return
    check(bat.exists(), "BAT 文件存在")
    if not bat.exists():
        return

    out = WORK / "sample_图片版.xlsx"

    def run_bat(*extra: str, stdin: str = "") -> str:
        result = subprocess.run(["cmd", "/c", bat.name, *extra],
                                capture_output=True, text=True, encoding="utf-8",
                                errors="replace", cwd=str(ROOT), env=child_env(),
                                input=stdin)
        return result.stdout

    if out.exists():
        out.unlink()

    dragging = run_bat(str(WORK / "sample.xlsx"))
    check("没有找到 Python" not in dragging, "BAT 能自己找到 Python")
    check("[错误]" not in dragging, "拖拽模式没有报错")
    check(out.exists(), "把表格拖到 BAT 图标上即可处理")

    interactive = run_bat(stdin=str(WORK / "sample.csv") + "\n")
    check("把表格文件拖到本窗口" in interactive, "无参数双击时进入交互式输入")
    check("完成：写入" in interactive, "粘贴路径后处理完成")

    wrong = run_bat(stdin="C:\\nope\\xx.xlsx\n" + str(WORK / "sample.csv") + "\n")
    check("找不到这个文件" in wrong, "路径写错时提示并允许重新输入")
    check("完成：写入" in wrong, "重新输入正确路径后继续处理")

    missing = run_bat("C:\\nope\\xx.xlsx")
    check("处理没有完成" in missing, "拖入不存在的文件时给出失败提示")


def check_output(path: Path, bad_rows: set[int], base: str, label: str) -> None:
    print(f"\n[{label}] 校验 {path.name}")
    wb = load_workbook(path)
    ws = wb.worksheets[0]

    check(ws.cell(1, 2).value == "主图链接", "原 URL 列表头未被改动")
    check(ws.cell(1, 3).value == "主图链接 图片", "图片列插在原 URL 列右侧")
    check(ws.cell(2, 2).value == f"{base}/normal.jpg", "原 URL 内容保留")
    check(int(ws.cell(2, 4).value) == 99, "原有数据列顺序未被打乱")

    for row in range(DATA_FIRST, DATA_LAST + 1):
        value = ws.cell(row, 3).value
        if row in bad_rows:
            check(value == "下载失败", f"第 {row} 行（预期失败）标记为下载失败")
        else:
            check(value is None, f"第 {row} 行图片单元格为空（已贴图）")

    check(abs(ws.row_dimensions[2].height - (SIZE * 0.75 + 3)) < 0.01,
          "行高按图片尺寸调整")
    check(abs(ws.column_dimensions["C"].width - (SIZE / 7 + 0.7)) < 0.01,
          "图片列宽按图片尺寸调整")

    expected = (DATA_LAST - DATA_FIRST + 1) - len(bad_rows)
    with zipfile.ZipFile(path) as z:
        media = [n for n in z.namelist()
                 if n.startswith("xl/media/") and not n.endswith("/")]
        check(len(media) == expected, f"内嵌图片 {len(media)} 张（预期 {expected}）")
        check(all(n.endswith(".jpeg") for n in media), "图片以 JPEG 存储（体积最优）")

        drawing = z.read("xl/drawings/drawing1.xml").decode("utf-8")
        check(drawing.count("<pic>") == expected, "每个锚点对应一张图片")
        anchors = sorted(int(m) for m in re.findall(
            r"<from><col>\d+</col><colOff>\d+</colOff><row>(\d+)</row>", drawing))
        expected_anchors = sorted(r - 1 for r in range(DATA_FIRST, DATA_LAST + 1)
                                  if r not in bad_rows)
        check(anchors == expected_anchors, f"图片锚定到正确行（0 基）{anchors}")
        cols = sorted({int(m) for m in re.findall(r"<from><col>(\d+)</col>", drawing)})
        check(cols == [2], f"图片全部锚定在图片列（第 3 列），实际 {cols}")

    check(path.stat().st_size < 500 * 1024, f"文件体积 {path.stat().st_size / 1024:.0f} KB 合理")


def check_thumb_rendering() -> None:
    print("\n[单元] 缩略图渲染")
    import url2img

    out = WORK / "thumb_transparent.jpg"
    url2img.build_thumb((IMG_DIR / "transparent.png").read_bytes(), SIZE, str(out))
    img = Image.open(out)
    check(img.size == (SIZE, SIZE), "输出为固定尺寸正方形")
    check(img.getpixel((3, 3))[0] > 240, "透明区域合成为白色（未发黑）")
    center = img.getpixel((50, 50))
    check(center[0] > 150 and center[1] < 120, f"图形内容保留在中心 {center}")
    check(out.stat().st_size < 20 * 1024, f"缩略图体积 {out.stat().st_size / 1024:.1f} KB")

    out2 = WORK / "thumb_wide.jpg"
    url2img.build_thumb((IMG_DIR / "wide.png").read_bytes(), SIZE, str(out2))
    wide = Image.open(out2)
    check(wide.size == (SIZE, SIZE), "横长图等比缩放后不超框")
    check(wide.getpixel((50, 4))[0] > 240, "横长图上下留白")


def check_url_extraction() -> None:
    print("\n[单元] URL 识别")
    import url2img

    cases = [
        ("http://a.com/x.jpg", ["http://a.com/x.jpg"]),
        ("商品图片：http://a.com/x.png 请查收", ["http://a.com/x.png"]),
        ("http://a.com/x.jpg,http://a.com/y.png",
         ["http://a.com/x.jpg", "http://a.com/y.png"]),
        ("http://a.com/x.webp?w=1&h=2", ["http://a.com/x.webp?w=1&h=2"]),
        ("https://img.360buyimg.com/n1/jfs/t1/a.jpg_430x430q90.jpg",
         ["https://img.360buyimg.com/n1/jfs/t1/a.jpg_430x430q90.jpg"]),
        ("http://a.com/x.pdf", []),
        ("没有链接", []),
        # 未转义的逗号属于 OSS 缩略参数，不能被当成链接分隔符而截断
        ("http://a.com/x.png?x-oss-process=image/resize,m_mfit,w_400,h_400",
         ["http://a.com/x.png?x-oss-process=image/resize,m_mfit,w_400,h_400"]),
        # 未转义参数 + 逗号分隔多图，两条都要完整取到
        ("http://a.com/p.png?process=image/resize,w_50,"
         "http://b.com/q.png?process=image/resize,w_50",
         ["http://a.com/p.png?process=image/resize,w_50",
          "http://b.com/q.png?process=image/resize,w_50"]),
    ]
    for text, expected in cases:
        got = url2img.find_urls(text)
        check(got == expected, f"{text[:46]!r} → {got}")

    # 真实电商后台导出的形态：JSON 里放同一张图的多种尺寸
    json_cell = (
        '{"pic400x400":"http://jkyun.oss-cn-hangzhou.aliyuncs.com/a.png?Expires=1'
        '&OSSAccessKeyId=key&Signature=abc%3D'
        '&x-oss-process=image%2Fresize%2Cm_mfit%2Cw_400%2Ch_400%2Fformat%2Cjpg",'
        '"pic0x0":"http://jkyun.oss-cn-hangzhou.aliyuncs.com/a.png?Expires=1'
        '&OSSAccessKeyId=key&Signature=def%3D",'
        '"pic50x50":"http://jkyun.oss-cn-hangzhou.aliyuncs.com/a.png?Expires=1'
        '&OSSAccessKeyId=key&Signature=ghi%3D'
        '&x-oss-process=image%2Fresize%2Cm_mfit%2Cw_50%2Ch_50%2Fformat%2Cjpg"}'
    )
    urls = url2img.find_urls(json_cell)
    check(len(urls) == 3, f"JSON 单元格里 3 个链接全部找到（实际 {len(urls)}）")
    check(bool(urls) and urls[0].endswith("%2Cjpg"),
          "取第一个链接，即 400x400 那版，且尾部未被截断")
    check(all(u.startswith("http") and "Signature=" in u for u in urls),
          "每个链接都保住了自己的签名参数")


def main() -> int:
    if WORK.exists():
        shutil.rmtree(WORK)
    WORK.mkdir(parents=True)

    make_images()
    server, port = start_server()
    base = f"http://127.0.0.1:{port}"
    print(f"本地图片服务：{base}")

    try:
        sample = make_sample(port)

        print("\n[运行 1] 默认参数（Referer 自动生成）")
        run_script(str(sample))
        out = WORK / "sample_图片版.xlsx"
        check_output(out, set(BAD_ROWS), base, "默认参数")

        print("\n[请求统计] 验证去重与重试")
        check(HITS.get("/normal.jpg") == 3,
              f"/normal.jpg 请求 {HITS.get('/normal.jpg')} 次（A/I/J 同一链接合并为 1 次，"
              f"H 的 ?w=200 与 K 的 x-oss-process 各 1 次）")
        check(HITS.get("/wide.png") == 1,
              f"/wide.png 请求 {HITS.get('/wide.png')} 次（同格多图只取第一张）")
        check(HITS.get("/missing.jpg") == 3,
              f"/missing.jpg 请求 {HITS.get('/missing.jpg')} 次（失败重试 2 次生效）")

        print("\n[运行 2] --referer none（防盗链应失败）")
        HITS.clear()
        run_script(str(sample), "--referer", "none", "-o", str(WORK / "noref.xlsx"))
        check_output(WORK / "noref.xlsx", {7, 8}, base, "无 Referer")
        check(HITS.get("/hotlink.jpg", 0) >= 1, "防盗链图片在无 Referer 时被拒（证明 Referer 机制有效）")

        print("\n[运行 3] CSV 输入")
        run_script(str(WORK / "sample.csv"), "-o", str(WORK / "csv_out.xlsx"))
        check_output(WORK / "csv_out.xlsx", set(BAD_ROWS), base, "CSV 输入")

        print("\n[运行 4] 指定列 --column 主图链接")
        run_script(str(sample), "--column", "主图链接", "-o", str(WORK / "col.xlsx"))
        check((WORK / "col.xlsx").exists(), "指定列模式正常输出")

        check_bat()
        check_thumb_rendering()
        check_url_extraction()
    finally:
        server.shutdown()

    print("\n" + "=" * 56)
    if FAILURES:
        print(f"失败 {len(FAILURES)} 项：")
        for item in FAILURES:
            print(f"  - {item}")
        return 1
    print("全部通过")
    return 0


if __name__ == "__main__":
    sys.exit(main())
